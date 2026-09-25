"""Export des avis en CSV, JSON et Excel."""

from __future__ import annotations

import csv
import io
import json

from .scraper import FIELDS


def to_csv(reviews: list[dict], delimiter: str = ",", fields: list[str] = FIELDS) -> bytes:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=fields, delimiter=delimiter, extrasaction="ignore")
    w.writeheader()
    w.writerows(reviews)
    # BOM UTF-8 : Excel ouvre ainsi correctement les accents
    return buf.getvalue().encode("utf-8-sig")


def to_json(reviews: list[dict], business: dict | None = None) -> bytes:
    payload = {"business": business or {}, "count": len(reviews), "reviews": reviews}
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")


def to_xlsx(reviews: list[dict], business: dict | None = None, fields: list[str] = FIELDS) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Avis"
    ws.append(fields)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="00B67A")
    for r in reviews:
        ws.append([r.get(f, "") for f in fields])

    widths = {"title": 40, "text": 80, "reply": 60, "url": 45, "date": 22, "name": 22}
    for i, f in enumerate(fields, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = widths.get(f, 12)
    long_cols = {i for i, f in enumerate(fields, start=1) if f in ("title", "text", "reply")}
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in long_cols)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    if business:
        info = wb.create_sheet("Infos")
        for k, v in business.items():
            info.append([k, v if not isinstance(v, (dict, list)) else json.dumps(v)])
        info.column_dimensions["A"].width = 18
        info.column_dimensions["B"].width = 50

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


EXPORTERS = {
    # (type MIME, fonction(avis, infos, colonnes) -> octets)
    "csv": ("text/csv; charset=utf-8", lambda r, b, f=FIELDS: to_csv(r, fields=f)),
    "json": ("application/json", lambda r, b, f=FIELDS: to_json(r, b)),
    "xlsx": (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        to_xlsx,
    ),
}
